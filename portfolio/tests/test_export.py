"""Exports must be readable without this program: DictReader for the CSVs,
openpyxl for the workbook, and every number as the text of a Decimal."""

from __future__ import annotations

import csv
import datetime as dt
import importlib.util
import io
import sys
from decimal import Decimal

import numpy as np
import pandas as pd
import pytest

from portfolio.core.models import AssetClass, Instrument, Transaction, TransactionType as T
from portfolio.core.money import FxTable, Money
from portfolio.core.positions import PriceQuote, derive_positions
from portfolio.core.report import holdings_table
from portfolio.data.export import (HOLDING_COLUMNS, TRANSACTION_COLUMNS, ExportUnavailable,
                                   holdings_csv, transactions_csv, workbook)

D = dt.date
A, B = "IE0002Y8CX98", "IE000IAXNM41"
TODAY = D(2026, 9, 3)


@pytest.fixture
def instruments():
    return {
        A: Instrument(A, "WisdomTree Europe Defence", AssetClass.ETF, "EUR"),
        B: Instrument(B, "iShares Europe Defence", AssetClass.ETF, "EUR"),
    }


@pytest.fixture
def ledger():
    return [
        Transaction(D(2025, 1, 10), B, T.BUY, Decimal("200"), Decimal("4.00"), "EUR",
                    Decimal("3.00"), note="second, but earlier in the file"),
        Transaction(D(2025, 1, 5), A, T.BUY, Decimal("100"), Decimal("10.00"), "EUR",
                    Decimal("5.00"), note='note with, comma and "quotes"'),
        Transaction(D(2025, 3, 1), A, T.SELL, Decimal("40"), Decimal("12.00"), "EUR",
                    Decimal("1.00")),
    ]


class TestTransactionsCsv:
    def test_round_trips_through_dict_reader(self, ledger, instruments):
        rows = list(csv.DictReader(io.StringIO(transactions_csv(ledger, instruments))))
        assert list(rows[0]) == TRANSACTION_COLUMNS
        assert len(rows) == 3
        first = rows[0]
        assert first["date"] == "2025-01-05" and first["isin"] == A
        assert first["name"] == "WisdomTree Europe Defence" and first["type"] == "BUY"
        assert Decimal(first["quantity"]) == 100 and Decimal(first["price_per_unit"]) == 10
        assert first["note"] == 'note with, comma and "quotes"'

    def test_sorted_by_date_then_id(self, ledger, instruments):
        rows = list(csv.DictReader(io.StringIO(transactions_csv(ledger, instruments))))
        assert [r["date"] for r in rows] == ["2025-01-05", "2025-01-10", "2025-03-01"]

    def test_gross_and_net_match_the_ledger(self, ledger, instruments):
        rows = list(csv.DictReader(io.StringIO(transactions_csv(ledger, instruments))))
        buy, sell = rows[0], rows[2]
        assert Decimal(buy["gross"]) == Decimal("1000.00")
        assert Decimal(buy["net"]) == Decimal("-1005.00"), "money out is negative"
        assert Decimal(sell["gross"]) == Decimal("480.00")
        assert Decimal(sell["net"]) == Decimal("479.00")

    def test_unknown_instrument_gets_a_blank_name_not_an_error(self, ledger):
        rows = list(csv.DictReader(io.StringIO(transactions_csv(ledger, {}))))
        assert all(r["name"] == "" for r in rows)

    def test_ids_are_preserved(self, ledger, instruments):
        rows = list(csv.DictReader(io.StringIO(transactions_csv(ledger, instruments))))
        assert {r["id"] for r in rows} == {t.id for t in ledger}

    def test_empty_ledger_is_a_header_only(self, instruments):
        text = transactions_csv([], instruments)
        assert text.strip() == ",".join(TRANSACTION_COLUMNS)


class TestHoldingsCsv:
    def _table(self, instruments, ledger, quotes, rates=None):
        pos = derive_positions(ledger, instruments, rates=rates, quotes=quotes)
        return holdings_table(pos, instruments, rates=rates, as_of=TODAY)

    def test_round_trips_through_dict_reader(self, instruments, ledger):
        quotes = {A: PriceQuote(Money(Decimal("12.00"), "EUR"), TODAY, "test", 15),
                  B: PriceQuote(Money(Decimal("4.00"), "EUR"), TODAY, "test", 15)}
        rows = list(csv.DictReader(io.StringIO(
            holdings_csv(self._table(instruments, ledger, quotes)))))
        assert list(rows[0]) == HOLDING_COLUMNS
        by_isin = {r["isin"]: r for r in rows}
        b = by_isin[B]
        assert b["name"] == "iShares Europe Defence"
        assert Decimal(b["quantity"]) == 200 and Decimal(b["price"]) == 4
        assert Decimal(b["value"]) == Decimal("800.00")
        assert Decimal(b["unrealised"]) == Decimal("-3.00")
        assert Decimal(b["weight"]) + Decimal(by_isin[A]["weight"]) == 1
        assert "delayed ~15m" in b["price_as_of"]

    def test_unpriced_row_has_blank_money_cells_and_a_warning(self, instruments, ledger):
        quotes = {A: PriceQuote(Money(Decimal("12.00"), "EUR"), TODAY, "test", 15)}
        rows = list(csv.DictReader(io.StringIO(
            holdings_csv(self._table(instruments, ledger, quotes)))))
        b = next(r for r in rows if r["isin"] == B)
        assert b["price"] == "" and b["value"] == "" and b["weight"] == ""
        assert "excluded from the total" in b["warnings"]

    def test_warnings_are_joined_and_fx_note_travels_with_the_price(self, instruments, ledger):
        fx = FxTable().add("USD", "EUR", TODAY, "0.90")
        quotes = {A: PriceQuote(Money(Decimal("12.00"), "USD"), D(2025, 9, 10), "test", None,
                                is_stale=True),
                  B: PriceQuote(Money(Decimal("4.00"), "EUR"), TODAY, "test", 15)}
        rows = list(csv.DictReader(io.StringIO(
            holdings_csv(self._table(instruments, ledger, quotes, rates=fx)))))
        a = next(r for r in rows if r["isin"] == A)
        assert "NOT a current price" in a["warnings"]
        assert "not a live quote" in a["warnings"]
        assert "USD->EUR at 0.9000" in a["warnings"]
        assert a["warnings"].count(" | ") == 2


# openpyxl ships with the app extra, not with core. The core CI job installs
# neither, on purpose, and the export must then fail with a sentence rather
# than a traceback -- which the last class below checks without openpyxl.
@pytest.mark.skipif(importlib.util.find_spec("openpyxl") is None,
                    reason="workbook export needs openpyxl (the app extra)")
class TestWorkbook:
    def _read(self, data: bytes):
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(data))

    def test_bytes_are_a_zip_and_sheets_read_back(self):
        frame = pd.DataFrame({"isin": [A, B], "value": [1200.5, 800.0]})
        data = workbook({"Holdings": frame, "Monthly": pd.DataFrame({"m": [1]})})
        assert data[:2] == b"PK"
        book = self._read(data)
        assert book.sheetnames == ["Holdings", "Monthly"]
        sheet = book["Holdings"]
        assert sheet["A1"].value == "isin" and sheet["B1"].value == "value"
        assert sheet["A2"].value == A and sheet["B3"].value == 800.0

    def test_header_is_bold_and_frozen(self):
        book = self._read(workbook({"S": pd.DataFrame({"a": [1]})}))
        sheet = book["S"]
        assert sheet["A1"].font.bold is True
        assert sheet.freeze_panes == "A2"

    def test_column_widths_follow_content(self):
        frame = pd.DataFrame({"short": [1], "long": ["x" * 40]})
        sheet = self._read(workbook({"S": frame}))["S"]
        assert sheet.column_dimensions["A"].width == 8
        assert sheet.column_dimensions["B"].width == 42

    def test_nan_becomes_an_empty_cell_not_the_text_nan(self):
        frame = pd.DataFrame({"a": [1.0, np.nan]})
        sheet = self._read(workbook({"S": frame}))["S"]
        assert sheet["A3"].value is None

    def test_meaningful_index_is_written_as_the_first_column(self):
        monthly = pd.DataFrame({"1": [0.01], "2": [-0.02]}, index=pd.Index([2025], name="year"))
        sheet = self._read(workbook({"Monthly": monthly}))["Monthly"]
        assert sheet["A1"].value == "year" and sheet["A2"].value == 2025
        assert sheet["B1"].value == "1" and sheet["C2"].value == -0.02

    def test_datetime_index_is_written_as_dates(self):
        frame = pd.DataFrame({"v": [1.0]}, index=pd.DatetimeIndex(["2025-01-15"]))
        sheet = self._read(workbook({"S": frame}))["S"]
        assert sheet["A2"].value.date() == D(2025, 1, 15)

    def test_decimals_and_numpy_scalars_are_numbers(self):
        frame = pd.DataFrame({"d": [Decimal("1.25")], "n": [np.int64(3)]})
        sheet = self._read(workbook({"S": frame}))["S"]
        assert sheet["A2"].value == 1.25 and sheet["B2"].value == 3

    def test_sheet_names_are_made_legal_for_excel(self):
        frames = {"Risk/Return: 2025?": pd.DataFrame({"a": [1]}),
                  "x" * 40: pd.DataFrame({"a": [1]}), ("x" * 40) + "y": pd.DataFrame({"a": [1]})}
        names = self._read(workbook(frames)).sheetnames
        assert names[0] == "Risk_Return_ 2025_"
        assert all(len(n) <= 31 for n in names) and len(set(names)) == 3

    def test_no_sheets_still_yields_a_valid_file(self):
        assert self._read(workbook({})).sheetnames == ["Sheet"]


class TestWorkbookWithoutOpenpyxl:
    def test_missing_openpyxl_is_a_sentence_with_the_install_command(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "openpyxl", None)
        with pytest.raises(ExportUnavailable) as exc:
            workbook({"S": pd.DataFrame({"a": [1]})})
        assert isinstance(exc.value, RuntimeError)
        assert "pip install" in str(exc.value)
