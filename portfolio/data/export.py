"""Exports: the ledger and the holdings as CSV, any set of frames as xlsx.

The data must outlive the tool. The CSVs here are readable by a spreadsheet
and by `csv.DictReader` with no help, and the transaction export carries
enough columns (gross, net, name) that a reader can reconcile it without
running this program. Every number is written as `Decimal` text, never as a
float, so what you paid is what the file says.

openpyxl is an `app` extra rather than a core dependency: `core/` must import
without it, and a user who never presses "download workbook" never needs it.
It is imported inside `workbook`, and its absence is reported as a sentence
with the install command rather than as an ImportError traceback.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import math
from decimal import Decimal
from typing import Any, Iterable

import pandas as pd

from ..core.models import Instrument, Transaction
from ..core.money import Money
from ..core.report import HoldingsTable

__all__ = ["transactions_csv", "holdings_csv", "workbook", "ExportUnavailable",
           "TRANSACTION_COLUMNS", "HOLDING_COLUMNS"]

TRANSACTION_COLUMNS = ["id", "date", "isin", "name", "type", "quantity",
                       "price_per_unit", "currency", "fees", "gross", "net", "note"]
HOLDING_COLUMNS = ["isin", "name", "quantity", "avg_cost", "price", "price_as_of",
                   "value", "unrealised", "unrealised_pct", "weight", "warnings"]


class ExportUnavailable(RuntimeError):
    """An export format whose library is not installed. Carries the fix."""


def _money(value: Money | None) -> str:
    return "" if value is None else str(value.amount)


def _decimal(value: Decimal | None) -> str:
    return "" if value is None else str(value)


def transactions_csv(transactions: Iterable[Transaction],
                     instruments: dict[str, Instrument] | Iterable[Instrument]) -> str:
    """The ledger, oldest first, with the instrument name beside the ISIN.

    `gross` is quantity x price and `net` the signed cash effect in the
    transaction's own currency, both as the ledger computes them -- so a
    row can be checked by hand against a broker statement.
    """
    names = ({k: v.name for k, v in instruments.items()}
             if isinstance(instruments, dict)
             else {i.isin: i.name for i in instruments})
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=TRANSACTION_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for txn in sorted(transactions, key=lambda t: (t.date, t.id)):
        writer.writerow({
            "id": txn.id, "date": txn.date.isoformat(), "isin": txn.isin,
            "name": names.get(txn.isin, ""), "type": txn.type.value,
            "quantity": str(txn.quantity), "price_per_unit": str(txn.price_per_unit),
            "currency": txn.currency, "fees": str(txn.fees),
            "gross": str(txn.gross.amount), "net": str(txn.cash_flow().amount),
            "note": txn.note,
        })
    return out.getvalue()


def holdings_csv(table: HoldingsTable) -> str:
    """The Holdings table as displayed, warnings included.

    An unpriced row is written with blank value cells rather than zeros, and
    its warning travels with it: a spreadsheet that sums the column gets the
    same total the page shows, and the reason for the gap is on the row.
    The FX note joins the warnings column so a price in another currency is
    never a bare number.
    """
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=HOLDING_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for row in table.rows:
        notes = list(row.warnings)
        if row.fx_note:
            notes.append(row.fx_note)
        writer.writerow({
            "isin": row.isin, "name": row.name, "quantity": str(row.quantity),
            "avg_cost": _money(row.average_cost), "price": _money(row.price),
            "price_as_of": row.price_as_of, "value": _money(row.market_value),
            "unrealised": _money(row.unrealised),
            "unrealised_pct": _decimal(row.unrealised_pct),
            "weight": _decimal(row.weight), "warnings": " | ".join(notes),
        })
    return out.getvalue()


# --------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------

_SHEET_FORBIDDEN = str.maketrans({c: "_" for c in "[]:*?/\\"})


def _sheet_name(name: str, used: set[str]) -> str:
    # Excel refuses names over 31 characters or containing []:*?/\ -- and
    # refuses the whole file, not the sheet, so it has to be fixed here.
    clean = (str(name).translate(_SHEET_FORBIDDEN).strip() or "Sheet")[:31]
    candidate, n = clean, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = clean[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _cell(value: Any) -> Any:
    """A value openpyxl will write as what it is, or None for a gap."""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        stamp = value.to_pydatetime()
        if stamp.tzinfo is not None:
            stamp = stamp.astimezone(dt.timezone.utc).replace(tzinfo=None)
        return stamp.date() if stamp.time() == dt.time() else stamp
    if isinstance(value, Money):
        return value.amount
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()                       # numpy scalar -> Python
        except (ValueError, AttributeError):
            pass
        if isinstance(value, float) and math.isnan(value):
            return None
    if isinstance(value, (bool, int, float, Decimal, str, dt.date, dt.datetime, dt.time)):
        return value
    return str(value)


def _frame_rows(frame: pd.DataFrame) -> tuple[list[str], list[list[Any]]]:
    """Header and rows, carrying a meaningful index as the first column."""
    include_index = not isinstance(frame.index, pd.RangeIndex)
    header: list[str] = []
    if include_index:
        header.append(str(frame.index.name or "index"))
    header.extend(str(c) for c in frame.columns)
    rows: list[list[Any]] = []
    for idx, values in zip(frame.index, frame.itertuples(index=False, name=None)):
        cells = [_cell(v) for v in values]
        if include_index:
            cells.insert(0, _cell(idx))
        rows.append(cells)
    return header, rows


def workbook(sheets: dict[str, pd.DataFrame]) -> bytes:
    """One sheet per frame: bold header, frozen header row, fitted widths."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ExportUnavailable(
            "the workbook export needs openpyxl, which is not installed. Run: "
            "pip install -e \".[app]\"") from exc

    book = Workbook()
    book.remove(book.active)
    used: set[str] = set()
    for name, frame in sheets.items():
        sheet = book.create_sheet(_sheet_name(name, used))
        header, rows = _frame_rows(frame)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for col, title in enumerate(header, start=1):
            longest = max([len(title)] + [len(_text(r[col - 1])) for r in rows])
            # Two characters of air, clamped so a long note cannot produce a
            # column wider than the screen or a bare number a column of 2.
            sheet.column_dimensions[get_column_letter(col)].width = min(60, max(8, longest + 2))
    if not book.sheetnames:
        book.create_sheet("Sheet")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)
